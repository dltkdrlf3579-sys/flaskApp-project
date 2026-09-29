"""Read-only Navigator: configured queries, scoped counts, no persisted snapshot."""
from collections import Counter
import configparser
from datetime import date, datetime, time, timezone, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import logging
from numbers import Number
from pathlib import Path
import re

from flask import Blueprint, jsonify, render_template, send_file, session
from permission_helpers import enforce_permission, get_user_permission_level

navigator_bp = Blueprint('navigator', __name__)
MENU_CODE = 'AX_EHS_NAVIGATOR'
CONFIG_PATH = Path(__file__).with_name('navigator_config.ini')
logger = logging.getLogger(__name__)
EXCEPTION_FIELDS = ('exception_id', 'partner_id', 'partner_name', 'work_name',
                    'issue', 'status', 'category', 'owner_name', 'owner_login_id', 'owner_dept_id')
EXCEPTION_PRIVATE_COLUMNS = {'owner_login_id', 'owner_dept_id'}
LOCAL_TIMEZONE = timezone(timedelta(hours=9))


def load_settings(path=None):
    config = configparser.ConfigParser(interpolation=None)
    config.optionxform = str
    with open(path or CONFIG_PATH, encoding='utf-8-sig') as stream:
        config.read_file(stream)
    source = config.get('NAVIGATOR', 'source')
    if source not in {'qadb', 'sample_postgres'}:
        raise ValueError('Unsupported Navigator source')
    statuses = []
    completed = {s.strip().casefold() for s in config.get('NAVIGATOR', 'completed_columns').split(',')}
    for name, color in config.items('STATUS_COLUMNS'):
        if not re.fullmatch(r'#[0-9a-fA-F]{6}', color.strip()):
            raise ValueError('Invalid status color')
        statuses.append({'key': name.casefold(), 'label': name, 'color': color.strip(),
                         'completed': name.casefold() in completed})
    keys = [s['key'] for s in statuses]
    reserved = {'partner_id', 'partner_name', 'site', 'tier', 'work_name', 'work_id', 'scope',
                'owner_login_id', 'owner_dept_id', 'gate_completed', 'gate_incomplete'}
    if (not statuses or len(keys) != len(set(keys)) or reserved.intersection(keys)
            or not completed or not completed.issubset(keys)):
        raise ValueError('Invalid status columns')
    queries = {name: config.get('QUERIES', name, fallback='').strip()
               for name in ('workload', 'exceptions')}
    works = []
    work_ids = set()
    for section in config.sections():
        if not section.startswith('WORK:'):
            continue
        work_id = section[len('WORK:'):].strip()
        name = config.get(section, 'name', fallback=work_id).strip()
        scope = config.get(section, 'scope').strip().casefold()
        enabled = config.getboolean(section, 'enabled', fallback=True)
        query = config.get(section, 'query', fallback='').strip()
        if not work_id or work_id.casefold() in work_ids or not name:
            raise ValueError('Missing or duplicate work ID/name')
        if scope not in {'company', 'site'}:
            raise ValueError('Work scope must be company or site')
        if enabled and not query:
            raise ValueError('Missing enabled work query')
        work_ids.add(work_id.casefold())
        works.append({'id': work_id, 'name': name, 'scope': scope,
                      'enabled': enabled, 'query': query})
    if works and queries['workload']:
        raise ValueError('Use WORK queries or legacy workload, not both')
    if not works and not queries['workload']:
        raise ValueError('Missing work queries')
    return {'source': source, 'statuses': statuses, 'queries': queries, 'works': works}


def query_rows(sql, source):
    # SQL is trusted deployment configuration, never supplied by HTTP parameters.
    if source == 'qadb':
        from database_config import execute_SQL
        frame = execute_SQL(sql)
        return frame.to_dict(orient='records'), list(frame.columns)
    from db_connection import get_db_connection
    conn = get_db_connection()
    try:
        conn.execute('SET TRANSACTION READ ONLY')
        conn.execute('SET LOCAL statement_timeout = 15000')
        with conn.cursor() as cursor:
            cursor.execute(sql)
            columns = [column[0] for column in cursor.description]
            return [dict(row) for row in cursor.fetchall()], columns
    finally:
        conn.rollback()
        conn.close()


def normalize_rows(rows, columns, required):
    names = [str(name).casefold() for name in columns]
    if len(names) != len(set(names)) or not set(required).issubset(names):
        raise ValueError('Missing or duplicate Navigator query columns')
    return [{str(k).casefold(): v for k, v in row.items()} for row in rows]


def text_value(value):
    if value is None:
        return ''
    # pandas may return NaN/NaT for missing IDs. They must never match a scope.
    if str(value).strip().lower() in {'nan', 'nat', '<na>'}:
        return ''
    return str(value).strip()


def scope_rows(rows, level, user_id, dept_id):
    if level >= 3:
        return rows
    if level <= 0:
        return []
    user_id, dept_id = text_value(user_id), text_value(dept_id)
    return [row for row in rows
            if (user_id and text_value(row.get('owner_login_id')) == user_id)
            or (level == 2 and dept_id and text_value(row.get('owner_dept_id')) == dept_id)]


def count_value(value):
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number < 0 or number != number.to_integral_value():
            raise ValueError('Invalid count')
        if number > 9007199254740991:
            raise ValueError('Count exceeds browser integer precision')
        return int(number)
    except (InvalidOperation, TypeError) as exc:
        raise ValueError('Invalid count') from exc


def read_exceptions(settings, level, user_id='', dept_id='', runner=None):
    """Shared raw query and row scope for the dashboard and Excel download."""
    sql = settings['queries']['exceptions']
    if not sql:
        return [], [], False
    rows, columns = (runner or query_rows)(sql, settings['source'])
    rows = scope_rows(normalize_rows(rows, columns, EXCEPTION_FIELDS), level, user_id, dept_id)
    seen = set()
    for row in rows:
        identity = text_value(row['exception_id'])
        if not identity or identity in seen:
            raise ValueError('Missing or duplicate exception ID')
        seen.add(identity)
    return rows, columns, True


def exception_workbook(rows, columns):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    public_columns = [(str(column).casefold(), str(column)) for column in columns
                      if str(column).casefold() not in EXCEPTION_PRIVATE_COLUMNS]
    if len(rows) > 1048575 or len(public_columns) > 16384:
        raise ValueError('Exception data exceeds Excel worksheet limits')
    book = Workbook(write_only=True)
    sheet = book.create_sheet('RawData')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(public_columns))}{len(rows) + 1}'
    header_fill = PatternFill('solid', fgColor='244675')
    header_font = Font(name='맑은 고딕', color='FFFFFF', bold=True)
    body_font = Font(name='맑은 고딕', size=11)
    alignment = Alignment(vertical='top', wrap_text=True)

    def cell_value(value, key):
        if value is None or str(value).strip().lower() in {'nan', 'nat', '<na>'}:
            return None
        if key in {'exception_id', 'partner_id'}:
            return text_value(value)
        if isinstance(value, datetime) and value.utcoffset() is not None:
            return value.astimezone(LOCAL_TIMEZONE).replace(tzinfo=None)
        if isinstance(value, time) and value.utcoffset() is not None:
            return value.isoformat()
        if isinstance(value, (dict, list, tuple)):
            return json.dumps(value, ensure_ascii=False, default=str)
        if isinstance(value, (str, Number, Decimal, date, time, timedelta)):
            return value
        return str(value)

    def excel_cell(value, header=False):
        if isinstance(value, str) and len(value) > 32767:
            raise ValueError('Exception text exceeds Excel cell limits')
        cell = WriteOnlyCell(sheet, value=value)
        # Raw strings must stay text, including values beginning with '='.
        if isinstance(value, str):
            cell.data_type = 's'
        cell.font = header_font if header else body_font
        cell.alignment = alignment
        if header:
            cell.fill = header_fill
        elif isinstance(value, datetime):
            cell.number_format = 'yyyy-mm-dd hh:mm:ss'
        elif isinstance(value, date):
            cell.number_format = 'yyyy-mm-dd'
        return cell

    for index, (key, label) in enumerate(public_columns, start=1):
        lengths = [len(label)] + [len(str(cell_value(row.get(key), key) or '')) for row in rows[:100]]
        width = max(15, min(60, max(lengths) * 1.5 + 2))
        if key == 'issue':
            width = max(width, 36)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.append([excel_cell(label, header=True) for _, label in public_columns])
    for row in rows:
        sheet.append([excel_cell(cell_value(row.get(key), key)) for key, _ in public_columns])
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def workload_rows(settings, level, user_id, dept_id, runner):
    """Read each enabled work independently; scope is configuration, never inferred."""
    keys = [s['key'] for s in settings['statuses']]
    required = ['partner_id', 'partner_name', 'tier', 'owner_login_id', 'owner_dept_id', *keys]
    works = settings.get('works', [])
    if works:
        for work in works:
            if not work['enabled']:
                continue
            rows, columns = runner(work['query'], settings['source'])
            needed = required + (['site'] if work['scope'] == 'site' else [])
            rows = normalize_rows(rows, columns, needed)
            names = {str(name).casefold() for name in columns}
            if ('gate_completed' in names) != ('gate_incomplete' in names):
                raise ValueError('Gate columns must be supplied together')
            for row in scope_rows(rows, level, user_id, dept_id):
                yield {**row, 'work_id': work['id'], 'work_name': work['name'],
                       'scope': work['scope'],
                       'site': text_value(row.get('site')) if work['scope'] == 'site' else '',
                       'gate_completed': row.get('gate_completed', 0),
                       'gate_incomplete': row.get('gate_incomplete', 0)}
    else:
        # Existing deployed INIs keep their original site-based data contract.
        rows, columns = runner(settings['queries']['workload'], settings['source'])
        rows = normalize_rows(rows, columns, required + ['site', 'work_name', 'gate_completed', 'gate_incomplete'])
        for row in scope_rows(rows, level, user_id, dept_id):
            yield {**row, 'work_id': text_value(row['work_name']), 'scope': 'site'}


def build_dashboard(settings, level, user_id='', dept_id='', runner=query_rows):
    statuses = settings['statuses']
    keys = [s['key'] for s in statuses]
    groups = {'work': {}, 'sites': {}, 'partners': {}}
    total = {key: 0 for key in keys}
    gate = {'completed': 0, 'incomplete': 0}
    seen = set()
    def new_entry(identity, name):
        return {'id': identity, 'name': name, 'counts': {key: 0 for key in keys},
                'gate_completed': 0, 'gate_incomplete': 0, 'sites': set(), 'tiers': set()}

    def add_counts(entry, counts, gc, gi):
        for key in keys:
            entry['counts'][key] += counts[key]
        entry['gate_completed'] += gc
        entry['gate_incomplete'] += gi

    for row in workload_rows(settings, level, user_id, dept_id, runner):
        partner_id, site, work_id, work_name = (text_value(row[k]) for k in (
            'partner_id', 'site', 'work_id', 'work_name'))
        identity = (work_id, partner_id, site, text_value(row['owner_login_id']), text_value(row['owner_dept_id']))
        if (not all(text_value(row[k]) for k in ('partner_id', 'partner_name', 'work_name', 'tier'))
                or (row['scope'] == 'site' and not site)):
            raise ValueError('Missing workload identity')
        if identity in seen:
            raise ValueError('Duplicate workload grain')
        seen.add(identity)
        counts = {key: count_value(row[key]) for key in keys}
        completed = sum(counts[s['key']] for s in statuses if s['completed'])
        gc, gi = count_value(row['gate_completed']), count_value(row['gate_incomplete'])
        if gc > completed or gi > sum(counts.values()) - completed:
            raise ValueError('Gate counts must be a subset of workload counts')
        gate['completed'] += gc
        gate['incomplete'] += gi
        for key in keys:
            total[key] += counts[key]
        group_keys = [('work', work_id, work_name), ('partners', partner_id, text_value(row['partner_name']))]
        if row['scope'] == 'site':
            group_keys.append(('sites', site, site))
        for group_name, group_key, group_label in group_keys:
            entry = groups[group_name].setdefault(group_key, new_entry(group_key, group_label))
            add_counts(entry, counts, gc, gi)
            if site:
                entry['sites'].add(site)
            entry['tiers'].add(text_value(row['tier']))
            if group_name == 'work':
                entry['scope'] = row['scope']
            if group_name == 'partners':
                # Company tasks appear once; site tasks retain their site in details.
                detail_key = (work_id, site)
                details = entry.setdefault('company_work' if row['scope'] == 'company' else 'site_work', {})
                detail = details.setdefault(detail_key, new_entry(work_id, work_name))
                detail['site'] = site
                add_counts(detail, counts, gc, gi)
    for mapping in groups.values():
        for entry in mapping.values():
            entry['sites'] = sorted(entry['sites'])
            entry['tiers'] = sorted(entry['tiers'])
            entry['total'] = sum(entry['counts'].values())
            if mapping is groups['partners']:
                for field in ('company_work', 'site_work'):
                    details = []
                    for detail in entry.get(field, {}).values():
                        details.append({key: detail[key] for key in ('id', 'name', 'site', 'counts')}
                                       | {'total': sum(detail['counts'].values())})
                    entry[field] = details
    erows, _, exceptions_linked = read_exceptions(settings, level, user_id, dept_id, runner)
    exceptions = [{**{key: text_value(row[key]) for key in EXCEPTION_FIELDS
                      if key not in EXCEPTION_PRIVATE_COLUMNS},
                   'site': text_value(row.get('site'))} for row in erows]
    return {
        'statuses': statuses, 'sample': settings['source'] == 'sample_postgres',
        'queried_at': datetime.now(LOCAL_TIMEZONE).strftime('%Y.%m.%d %H:%M'),
        'summary': {'partners': len(groups['partners']), 'total': sum(total.values()), 'counts': total},
        'gate': gate, 'work': list(groups['work'].values()), 'sites': list(groups['sites'].values()),
        'partners': list(groups['partners'].values()), 'exceptions': exceptions,
        'exceptions_linked': exceptions_linked,
        'site_work_linked': (any(w['enabled'] and w['scope'] == 'site' for w in settings.get('works', []))
                             if settings.get('works') else True),
        'categories': dict(Counter(e['category'] for e in exceptions)),
    }


@navigator_bp.get('/ax-ehs-navigator')
def dashboard():
    denied = enforce_permission(MENU_CODE, 'view')
    if denied is not None:
        return denied
    from audit_logger import record_menu_view
    record_menu_view(MENU_CODE)
    return render_template('navigator.html')


@navigator_bp.get('/api/ax-ehs-navigator/dashboard')
def dashboard_data():
    denied = enforce_permission(MENU_CODE, 'view', response_type='json')
    if denied is not None:
        return denied
    try:
        payload = build_dashboard(load_settings(), get_user_permission_level(MENU_CODE, 'read'),
                                  session.get('user_id'), session.get('deptid'))
        response = jsonify(payload)
        response.headers['Cache-Control'] = 'no-store'
        return response
    except Exception:
        logger.exception('Navigator dashboard query failed')
        return jsonify({'error': '현황을 불러오지 못했습니다. 조회 설정과 데이터 연결을 확인해 주세요.'}), 503


@navigator_bp.get('/api/ax-ehs-navigator/exceptions/export')
def export_exceptions():
    denied = enforce_permission(MENU_CODE, 'view', response_type='json')
    if denied is not None:
        return denied
    try:
        rows, columns, linked = read_exceptions(
            load_settings(), get_user_permission_level(MENU_CODE, 'read'),
            session.get('user_id'), session.get('deptid'))
        if not linked:
            return jsonify({'error': '비정상 진행 현황 조회가 아직 연계되지 않았습니다.'}), 404
        response = send_file(
            exception_workbook(rows, columns), as_attachment=True,
            download_name=f'navigator_exceptions_{datetime.now(LOCAL_TIMEZONE):%Y%m%d_%H%M%S}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            max_age=0, etag=False)
        response.headers['Cache-Control'] = 'no-store, private'
        return response
    except Exception:
        logger.exception('Navigator exception Excel export failed')
        return jsonify({'error': '엑셀을 내려받지 못했습니다. 조회 설정과 데이터 연결을 확인해 주세요.'}), 503
